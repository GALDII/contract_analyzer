# core/document_parser.py
import docx
from io import BytesIO

def parse_document(uploaded_file) -> str:
    """
    Parses a DOCX file and extracts text content.
    
    Args:
        uploaded_file: Streamlit uploaded file object
        
    Returns:
        str: Extracted text from the document
    """
    try:
        # Read the uploaded file
        doc = docx.Document(BytesIO(uploaded_file.read()))
        
        # Extract text from all paragraphs
        full_text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():  # Only add non-empty paragraphs
                full_text.append(paragraph.text.strip())
        
        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    full_text.append(" | ".join(row_text))
        
        # Join all text with newlines
        contract_text = '\n'.join(full_text)
        
        if not contract_text.strip():
            raise ValueError("No text content found in the document")
            
        return contract_text
        
    except Exception as e:
        raise Exception(f"Error parsing document: {str(e)}")